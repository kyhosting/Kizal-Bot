import os
import re
from pyrogram import Client, filters
from pyrogram.types import Message
from openpyxl import load_workbook

from config import is_owner
from core.database import db
from core.limit_manager import limit_manager
from core.error_handler import handle_errors
from utils.keyboards import get_cancel_keyboard, get_menu_keyboard, get_limit_upgrade_inline
from utils.messages import Messages
from utils.file_utils import FileUtils
from utils.vcf_utils import VCFUtils
from utils.session_manager import SessionManager


def extract_numbers_from_excel(filepath: str) -> list:
    numbers = []
    try:
        wb = load_workbook(filepath, read_only=True)
        for sheet in wb.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value:
                        cell_str = str(cell.value)
                        found = re.findall(r'[\+]?[\d\s\-\(\)]+', cell_str)
                        for num in found:
                            clean_num = re.sub(r'[\s\-\(\)]', '', num)
                            if len(clean_num) >= 8:
                                numbers.append(clean_num)
        wb.close()
    except Exception:
        pass
    return numbers


@handle_errors
async def xlsx_to_vcf_start(client: Client, message: Message):
    user_id = message.from_user.id

    can_proceed, status = await limit_manager.check_limit(user_id)
    
    if not can_proceed:
        await message.reply_text(
            status["message"],
            reply_markup=get_limit_upgrade_inline()
        )
        return

    await SessionManager.save(user_id, "xlsx_to_vcf", step=1)

    await message.reply_text(
        Messages.xls_to_vcf_start(),
        reply_markup=get_cancel_keyboard()
    )


@handle_errors
async def handle_xlsx_to_vcf(client: Client, message: Message):
    user_id = message.from_user.id

    if message.text == "❌ BATAL ❌":
        await SessionManager.clear(user_id)
        FileUtils.cleanup_user_files(user_id)
        await message.reply_text(
            Messages.cancelled(),
            reply_markup=get_menu_keyboard(user_id)
        )
        return

    session = await SessionManager.get(user_id)
    if not session or session.get("mode") != "xlsx_to_vcf":
        return

    step = session.get("step", 0)
    data = session.get("data") or {}

    if step == 1:
        if not message.document:
            await message.reply_text("```\n❌ Kirim file .xlsx atau .xls!\n```")
            return

        filename = message.document.file_name
        if not (filename.endswith('.xlsx') or filename.endswith('.xls')):
            await message.reply_text("```\n❌ File harus berformat .xlsx atau .xls!\n```")
            return

        filepath = FileUtils.get_temp_path(user_id, filename)
        await message.download(file_name=filepath)

        numbers = extract_numbers_from_excel(filepath)

        if not numbers:
            FileUtils.safe_remove(filepath)
            await message.reply_text(
                "```\n❌ Tidak ada nomor ditemukan!\n```",
                reply_markup=get_menu_keyboard(user_id)
            )
            await SessionManager.clear(user_id)
            return

        data["phone_numbers"] = numbers
        data["xlsx_filepath"] = filepath

        await SessionManager.save(user_id, "xlsx_to_vcf", 2, data)
        await message.reply_text(
            f"```\n📝 NAMA FILE VCF\n───────────────────────────────────────\n\nTotal nomor ditemukan: {len(numbers)}\n\nMasukkan nama file output\n(tanpa ekstensi .vcf)\n\n───────────────────────────────────────\n```",
            reply_markup=get_cancel_keyboard()
        )

    elif step == 2:
        data["vcf_filename"] = message.text.strip()
        await SessionManager.save(user_id, "xlsx_to_vcf", 3, data)
        await message.reply_text(
            "```\n👤 NAMA KONTAK\n───────────────────────────────────────\n\nMasukkan format nama kontak\n\nContoh: kontak\nHasil: kontak 0001, kontak 0002, ...\n\n───────────────────────────────────────\n```",
            reply_markup=get_cancel_keyboard()
        )

    elif step == 3:
        contact_name = message.text.strip()
        phone_numbers = data.get("phone_numbers", [])
        vcf_filename = data.get("vcf_filename", "output")

        vcf_filepath = FileUtils.get_temp_path(user_id, f"{vcf_filename}.vcf")
        keyboard = get_menu_keyboard(user_id)

        try:
            VCFUtils.create_vcf_file(phone_numbers, contact_name, vcf_filepath)

            await message.reply_document(
                document=vcf_filepath,
                file_name=f"{vcf_filename}.vcf",
                caption=f"✅ Berhasil convert XLS to VCF!\n📂 Total: {len(phone_numbers)} kontak"
            )

            success, status = await limit_manager.use_operation(user_id, "xlsx_to_vcf")
            
            if status.get("message") and status.get("warning_level") in ["low", "moderate"]:
                await message.reply_text(status["message"], reply_markup=keyboard)

        except Exception as e:
            await message.reply_text(f"```\n❌ Error: {str(e)}\n```", reply_markup=keyboard)
        finally:
            FileUtils.safe_remove(data.get("xlsx_filepath", ""))
            FileUtils.safe_remove(vcf_filepath)
            await SessionManager.clear(user_id)


def register_xlsx_to_vcf_handlers(app: Client):
    app.on_message(filters.regex("^🜲 XLS TO VCF 🜲$") & filters.private)(xlsx_to_vcf_start)
